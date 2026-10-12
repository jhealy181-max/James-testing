"""Build University_experiential_learning_detailed_review.xlsx from per-institution JSON results.

Usage: python3 -I el_review/build.py
Reads   : University_learning_comparison (3).xlsx (never modified), el_review/institutions.json,
          el_review/legacy/*.json, el_review/results/*.json
Writes  : University_experiential_learning_detailed_review.xlsx,
          el_review/University_experiential_learning_institution_profiles.md,
          el_review/CHECKPOINT.md, el_review/qc_report.txt
"""
import json, os, re, glob, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ELR = os.path.join(ROOT, "el_review")
SRC = os.path.join(ROOT, "University_learning_comparison (3).xlsx")
OUT = os.path.join(ROOT, "University_experiential_learning_detailed_review.xlsx")
PROFILES = os.path.join(ELR, "University_experiential_learning_institution_profiles.md")
TODAY = "2026-10-12"

NAVY, SLATE, INK = "FF203B60", "FF526277", "FF172B4D"
F_TITLE = Font(name="Arial", size=15, bold=True, color=NAVY)
F_SUB = Font(name="Arial", size=10, color=SLATE)
F_HDR = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
F_BODY = Font(name="Arial", size=10, color=INK)
F_BOLD = Font(name="Arial", size=10, bold=True, color=INK)
F_H2 = Font(name="Arial", size=12, bold=True, color=NAVY)
F_KPI = Font(name="Arial", size=20, bold=True, color=NAVY)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_KPI = PatternFill("solid", fgColor="FFEEF2F7")
WRAP = Alignment(wrap_text=True, vertical="top")
FILL_YES = PatternFill("solid", fgColor="FFD9EAD3")
FILL_NO = PatternFill("solid", fgColor="FFF4CCCC")
FILL_UNC = PatternFill("solid", fgColor="FFFFF2CC")

REACHES = ["University-wide", "Cross-faculty", "Faculty/school", "Programme/course", "Unclear"]
TRI = ["Yes", "No", "Unclear"]
STATUSES = ["Current delivery confirmed", "Announced/forthcoming", "Historical delivery only",
            "Paused/discontinued", "Current status unclear"]
PARTS = ["Compulsory for all students", "Compulsory within specified programmes", "Optional/elective",
         "Selective/competitive", "Unclear"]
PTYPES = ["Business/commercial", "Public sector", "Charity/community", "Mixed", "Other", "Unclear", "None"]
AREAS = ["central", "business", "engineering_computing", "science", "arts_humanities_social_sciences", "other"]
AREA_LABEL = {"central": "Central provision", "business": "Business/management",
              "engineering_computing": "Engineering/computing", "science": "Science",
              "arts_humanities_social_sciences": "Arts/humanities/social sciences",
              "other": "Other (law, health, design, architecture…)"}
REGION = {"RG": "UK – Russell Group", "UK": "UK – additional", "AU": "Australia", "US": "US", "EU": "EU"}
REGION_ORDER = ["UK – Russell Group", "UK – additional", "Australia", "US", "EU"]

def s(v, n=32000):
    if v is None: return ""
    if isinstance(v, (list, tuple)): v = "; ".join(str(x) for x in v if x not in (None, ""))
    v = str(v).strip()
    return v[:n]

def norm_tri(v):
    v = s(v).capitalize()
    return v if v in TRI else "Unclear"

def norm_choice(v, opts, default):
    v = s(v)
    for o in opts:
        if v.lower() == o.lower(): return o
    for o in opts:
        if v and v.lower().startswith(o.lower()[:12]): return o
    return default

def classify(a, b, c):
    if "No" in (a, b, c): return "Outside strict definition"
    if (a, b, c) == ("Yes", "Yes", "Yes"): return "Meets all three"
    return "Potentially relevant; evidence incomplete"

def toks(x):
    return set(w for w in re.findall(r"[a-z0-9]+", s(x).lower()) if len(w) > 2 and w not in {"the", "and", "for", "programme", "program", "scheme"})

# ---------------------------------------------------------------- load data
insts = json.load(open(os.path.join(ELR, "institutions.json")))
legacy = {i["inst_id"]: json.load(open(os.path.join(ELR, "legacy", f"{i['inst_id']}.json"))) for i in insts}
results = {}
for f in glob.glob(os.path.join(ELR, "results", "*.json")):
    iid = os.path.basename(f)[:-5]
    try:
        results[iid] = json.load(open(f))
    except Exception as e:
        print("BAD JSON", f, e)

qc = []  # (severity, inst, item, message)
schemes, sources, exclusions, searchlog, summary = [], [], [], [], []
legacy_disp = {}  # legacy_row -> disposition text

for inst in insts:
    iid = inst["inst_id"]; reg = REGION[iid[:2]]
    r = results.get(iid)
    if not r:
        summary.append(dict(inst=inst, region=reg, status="Not started", r=None))
        for lr in legacy[iid]["legacy_rows"]:
            legacy_disp[lr["legacy_row"]] = "Not yet reassessed (institution not reviewed)"
        continue
    status = "Complete" if s(r.get("review_status")).lower().startswith("complete") else "Partial"
    # ---- search log
    sl = r.get("search_log") or {}
    for a in AREAS:
        e = sl.get(a) or {}
        st = s(e.get("status")) or "Not recorded"
        searchlog.append([iid, inst["institution"], AREA_LABEL[a], st, s(e.get("queries"), 3000),
                          s(e.get("key_sources"), 3000), s(e.get("finding"), 3000)])
        if st == "Not recorded" or st.lower().startswith("not searched"):
            qc.append(("WARN", iid, a, f"search area status: {st}"))
            if status == "Complete": status = "Partial"
    # ---- schemes
    local2id = {}
    raw = r.get("schemes") or []
    for k, sc in enumerate(raw, 1):
        local2id[s(sc.get("local_id")) or f"S{k}"] = f"{iid}-S{k:02d}"
    children = collections.Counter(s(sc.get("parent_local_id")) for sc in raw if sc.get("parent_local_id"))
    for k, sc in enumerate(raw, 1):
        sid = f"{iid}-S{k:02d}"; lid = s(sc.get("local_id")) or f"S{k}"
        a, b, c = norm_tri(sc.get("project_based")), norm_tri(sc.get("group_working")), norm_tri(sc.get("employer_connected"))
        reach = norm_choice(sc.get("academic_reach"), REACHES, "Unclear")
        ost = norm_choice(sc.get("operating_status"), STATUSES, "Current status unclear")
        part = norm_choice(sc.get("participation_requirement"), PARTS, "Unclear")
        ptype = norm_choice(sc.get("partner_type"), PTYPES, "Unclear")
        umb = bool(sc.get("is_umbrella")) and children.get(lid, 0) > 0
        parent = local2id.get(s(sc.get("parent_local_id")), "") if sc.get("parent_local_id") else ""
        srcids = []
        crit_support = collections.defaultdict(int)
        for j, so in enumerate(sc.get("sources") or [], 1):
            rid = f"{sid}-R{j}"; srcids.append(rid)
            sup = so.get("supports") or []
            if isinstance(sup, str): sup = [x.strip() for x in sup.split(",")]
            for x in sup: crit_support[s(x).lower()] += 1
            sources.append([rid, iid, sid, s(sc.get("official_name")), s(sup), s(so.get("url")), s(so.get("title")),
                            s(so.get("section")), s(so.get("pub_year")), s(so.get("accessed")) or TODAY,
                            s(so.get("retrieval")) or "search-index extract", s(so.get("extract"), 2000)])
            if not s(so.get("url")).startswith("http"):
                qc.append(("ERROR", iid, sid, f"source {j} has no valid URL"))
        cls = classify(a, b, c)
        counted = "No (umbrella; children counted)" if umb else "Yes"
        qualifying = (cls == "Meets all three" and ost == "Current delivery confirmed" and counted == "Yes")
        # QC
        for crit, val, ev, key in (("A project", a, sc.get("project_evidence"), "project"),
                                   ("B group", b, sc.get("group_evidence"), "group"),
                                   ("C employer", c, sc.get("employer_evidence"), "employer")):
            if val == "Yes" and (not s(ev) or crit_support.get(key, 0) == 0):
                qc.append(("ERROR" if not s(ev) else "WARN", iid, sid, f"{crit}=Yes but evidence text/source tag missing"))
        if not srcids: qc.append(("ERROR", iid, sid, "scheme has no sources"))
        if qualifying and not s(sc.get("latest_evidenced_year")):
            qc.append(("WARN", iid, sid, "current qualifying scheme lacks latest evidenced year"))
        if qualifying and crit_support.get("status", 0) == 0:
            qc.append(("WARN", iid, sid, "current qualifying scheme has no source tagged 'status'"))
        if reach == "University-wide" and crit_support.get("reach", 0) + crit_support.get("eligibility", 0) == 0:
            qc.append(("WARN", iid, sid, "University-wide reach without source tagged reach/eligibility"))
        if part == "Compulsory for all students" and reach != "University-wide":
            qc.append(("ERROR", iid, sid, "compulsory for all students but reach not University-wide"))
        if c == "Yes" and ptype in ("None", "Unclear"):
            qc.append(("WARN", iid, sid, f"employer-connected Yes but partner type {ptype}"))
        if b == "Yes" and re.search(r"\bindividual(ly)?\b", s(sc.get("group_evidence")).lower()) and "team" not in s(sc.get("group_evidence")).lower():
            qc.append(("WARN", iid, sid, "group=Yes but group evidence mentions individual work"))
        if re.search(r"simulat|virtual work experience|forage", (s(sc.get("description")) + s(sc.get("employer_evidence"))).lower()) and c == "Yes":
            qc.append(("WARN", iid, sid, "possible simulation marked employer-connected — check"))
        schemes.append(dict(iid=iid, inst=inst["institution"], country=inst["country"], region=reg, sid=sid,
            parent=parent, umb="Yes" if umb else "No", counted=counted, sc=sc, a=a, b=b, c=c, reach=reach,
            ost=ost, part=part, ptype=ptype, cls=cls, qualifying=qualifying, srcids=srcids))
        if s(sc.get("legacy_ref")):
            lr_match = None; best = 0
            for lr in legacy[iid]["legacy_rows"]:
                ov = len(toks(lr["Scheme / framework"]) & toks(sc.get("legacy_ref")))
                if ov > best: best, lr_match = ov, lr
            if lr_match: legacy_disp.setdefault(lr_match["legacy_row"], f"Reassessed as {sid} ({cls})")
    # ---- exclusions
    for k, ex in enumerate(r.get("exclusions") or [], 1):
        xid = f"{iid}-X{k:02d}"; lrow = ""
        if ex.get("legacy"):
            best = 0
            for lr in legacy[iid]["legacy_rows"]:
                ov = len(toks(lr["Scheme / framework"]) & toks(ex.get("name")))
                if ov > best: best, lrow = ov, lr["legacy_row"]
            if lrow: legacy_disp.setdefault(lrow, f"Excluded as {xid}: {s(ex.get('category'))}")
        exclusions.append([xid, iid, inst["institution"], reg, s(ex.get("name")), "Yes" if ex.get("legacy") else "No",
                           lrow, s(ex.get("category")), s(ex.get("reason"), 3000), s(ex.get("source_url")), s(ex.get("status_note"))])
    for lr in legacy[iid]["legacy_rows"]:
        if lr["legacy_row"] not in legacy_disp:
            legacy_disp[lr["legacy_row"]] = "NOT TRACED — check"
            qc.append(("ERROR", iid, f"legacy row {lr['legacy_row']}", f"legacy '{lr['Scheme / framework']}' not traced"))
    summary.append(dict(inst=inst, region=reg, status=status, r=r))

# ---------------------------------------------------------------- workbook
wb = openpyxl.load_workbook(SRC)
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)

dash = wb["Dashboard"]  # its two charts are preserved by openpyxl on load/save (verified in verify.py)

# Legacy schemes tab: rename, add disposition column, repoint formulas
old = wb["Schemes"]; old.title = "Legacy schemes (v3)"
LEG = "'Legacy schemes (v3)'!"
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("=") and "Schemes!" in cell.value:
                cell.value = cell.value.replace("Schemes!", LEG)
old["A2"] = "LEGACY (v3, research date 8 Oct 2026): experiential learning scheme/framework examples — superseded"
old["A3"] = ("Retained unchanged for traceability. These rows were NOT tested against the strict three-criterion definition. "
             "Column Q records how each row was handled in the detailed review (scheme register ID or exclusion ID). Use 'EL scheme register'.")
old.cell(6, 17, "Disposition in detailed review (Oct 2026)")
old.cell(6, 17).font = F_HDR; old.cell(6, 17).fill = FILL_HDR; old.cell(6, 17).alignment = WRAP
for rr in range(7, 159):
    old.cell(rr, 17, legacy_disp.get(rr, "")).font = F_BODY
    old.cell(rr, 17).alignment = WRAP
old.column_dimensions["Q"].width = 48
old.tables["SchemesTable"].ref = "A6:Q158"
from openpyxl.worksheet.table import TableColumn
t = old.tables["SchemesTable"]
if len(t.tableColumns) == 16:
    t.tableColumns.append(TableColumn(id=17, name="Disposition in detailed review (Oct 2026)"))

ix = wb["Institution index"]
ix["A4"] = (ix["A4"].value or "") + " NOTE (Oct 2026): columns I–M count LEGACY scheme rows; see 'EL institution summary' for the strict-definition review."
ix.cell(6, 14, "EL review ID"); ix.cell(6, 15, "Sample / QS note (Oct 2026 reconciliation)")
for cc in (14, 15):
    ix.cell(6, cc).font = F_HDR; ix.cell(6, cc).fill = FILL_HDR; ix.cell(6, cc).alignment = WRAP
QSNOTE = {
    "Adelaide University": "Merged institution (University of Adelaide + UniSA) operating from 2026; pre-2026 scheme evidence relates to predecessor universities.",
    "City St George’s": "Merged institution (City + St George's, 2024); evidence may be badged City or Bayes Business School.",
    "Ohio State": "QS source is an institution profile page rather than the ranking table; rank 201 sits at the edge of the published table.",
}
for i in insts:
    rr = i["index_row"]
    ix.cell(rr, 14, i["inst_id"]).font = F_BODY
    note = QSNOTE.get(i["institution"], "")
    if "universities/" in s(i["qs_source"]) and not note:
        note = "QS rank sourced from institution profile page (not ranking table); not re-verified in this review."
    ix.cell(rr, 15, note).font = F_BODY; ix.cell(rr, 15).alignment = WRAP
ix.column_dimensions["N"].width = 12; ix.column_dimensions["O"].width = 50
t = ix.tables["InstitutionindexTable"]; t.ref = "A6:O130"
if len(t.tableColumns) == 13:
    t.tableColumns.append(TableColumn(id=14, name="EL review ID"))
    t.tableColumns.append(TableColumn(id=15, name="Sample / QS note (Oct 2026 reconciliation)"))

def header(ws, title, sub, sub2, cols, widths, row=6):
    ws["A2"] = title; ws["A2"].font = F_TITLE
    ws["A3"] = sub; ws["A3"].font = F_SUB
    ws["A4"] = sub2; ws["A4"].font = F_SUB
    for j, (cname, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(row, j, cname); c.font = F_HDR; c.fill = FILL_HDR; c.alignment = WRAP
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = ws.cell(row + 1, 3)
    ws.sheet_view.showGridLines = False

def add_table(ws, name, ncols, nrows, row=6):
    if nrows == 0: return
    tb = Table(displayName=name, ref=f"A{row}:{L(ncols)}{row + nrows}")
    tb.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tb)

def tri_cf(ws, rng):
    first = rng.split(":")[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="Yes"'], fill=FILL_YES))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="No"'], fill=FILL_NO))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="Unclear"'], fill=FILL_UNC))

def dv(ws, rng, opts):
    d = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
    ws.add_data_validation(d); d.add(rng)

# ---------------- Scheme register
REG_COLS = [("Institution ID", 11), ("Institution", 22), ("Academic reach", 17), ("Scheme ID", 12), ("Parent ID", 12),
    ("Umbrella with counted children?", 11), ("Official scheme name", 34), ("English name / translation", 26), ("Module / subject code", 16),
    ("What students actually do", 70), ("Administrative host", 26), ("Eligible faculties / programmes", 34), ("Faculty areas", 20),
    ("Student level", 14), ("Year of study", 14), ("Eligibility restrictions", 30), ("Participation requirement", 20),
    ("Curricular status", 13), ("Credit and assessment", 32), ("Duration", 16), ("Delivery format", 18),
    ("A. Project-based", 11), ("A. Evidence", 40), ("B. Group working", 11), ("B. Evidence", 40),
    ("C. Employer / external organisation", 12), ("C. Evidence: how the organisation participates", 46),
    ("External partner type", 17), ("Partner role", 30), ("Named partners (evidenced)", 30), ("Student deliverable", 28),
    ("Qualification classification", 22), ("Operating status", 20), ("Latest evidenced delivery year", 12),
    ("Counts as distinct scheme", 12), ("Current & meets all three (counted)", 13), ("Published participation data", 30),
    ("Source IDs (see EL sources)", 24), ("Unresolved points", 40), ("Legacy row reassessed", 24), ("Country", 12), ("Region group", 18)]
ws = wb.create_sheet("EL scheme register", 1)
header(ws, "Experiential learning: scheme register (strict three-criterion test)",
       "One row per distinct scheme or materially different participation route. Criteria A/B/C are three-state: Yes = explicit evidence; No = evidence shows not met; Unclear = insufficient evidence.",
       "Column C = academic reach (controlled). Column AF is calculated from A/B/C. Column AJ = counted scheme (not an umbrella), meets all three and current delivery confirmed. Evidence retrieved via search-index extracts (see EL method & QC).",
       [c for c, _ in REG_COLS], [w for _, w in REG_COLS])
R0 = 7
for n, x in enumerate(sorted(schemes, key=lambda z: (z["iid"], z["sid"]))):
    sc = x["sc"]; rr = R0 + n
    vals = [x["iid"], x["inst"], x["reach"], x["sid"], x["parent"], x["umb"], s(sc.get("official_name")), s(sc.get("english_name")),
            s(sc.get("module_code")), s(sc.get("description"), 3000), s(sc.get("admin_host")), s(sc.get("eligible_faculties_programmes")),
            s([AREA_LABEL.get(a, a) for a in (sc.get("faculty_areas") or [])]), s(sc.get("student_level")), s(sc.get("year_of_study")),
            s(sc.get("eligibility_restrictions")), x["part"], s(sc.get("curricular_status")), s(sc.get("credit_assessment")),
            s(sc.get("duration")), s(sc.get("delivery_format")), x["a"], s(sc.get("project_evidence"), 2000), x["b"],
            s(sc.get("group_evidence"), 2000), x["c"], s(sc.get("employer_evidence"), 2000), x["ptype"], s(sc.get("partner_role")),
            s(sc.get("named_partners")), s(sc.get("deliverable")),
            f'=IF(OR(V{rr}="No",X{rr}="No",Z{rr}="No"),"Outside strict definition",IF(AND(V{rr}="Yes",X{rr}="Yes",Z{rr}="Yes"),"Meets all three","Potentially relevant; evidence incomplete"))',
            x["ost"], s(sc.get("latest_evidenced_year")), "Yes" if x["counted"] == "Yes" else "No",
            f'=IF(AND(AI{rr}="Yes",AF{rr}="Meets all three",AG{rr}="Current delivery confirmed"),"Yes","No")',
            s(sc.get("participation_data")), s(x["srcids"]), s(sc.get("unresolved"), 2000), s(sc.get("legacy_ref")), x["country"], x["region"]]
    for j, v in enumerate(vals, 1):
        c = ws.cell(rr, j, v); c.font = F_BODY; c.alignment = WRAP
    ws.cell(rr, 7).font = F_BOLD
NREG = len(schemes); RLAST = R0 + max(NREG, 1) - 1
add_table(ws, "SchemeRegister", len(REG_COLS), NREG)
if NREG:
    for col in ("V", "X", "Z"):
        tri_cf(ws, f"{col}{R0}:{col}{RLAST}"); dv(ws, f"{col}{R0}:{col}{RLAST}", TRI)
    dv(ws, f"C{R0}:C{RLAST}", REACHES); dv(ws, f"AG{R0}:AG{RLAST}", STATUSES); dv(ws, f"Q{R0}:Q{RLAST}", PARTS); dv(ws, f"AB{R0}:AB{RLAST}", PTYPES)
    ws.conditional_formatting.add(f"AF{R0}:AF{RLAST}", FormulaRule(formula=[f'AF{R0}="Meets all three"'], fill=FILL_YES))
    ws.conditional_formatting.add(f"AF{R0}:AF{RLAST}", FormulaRule(formula=[f'AF{R0}="Outside strict definition"'], fill=FILL_NO))
    ws.conditional_formatting.add(f"AF{R0}:AF{RLAST}", FormulaRule(formula=[f'LEFT(AF{R0},11)="Potentially"'], fill=FILL_UNC))
    ws.conditional_formatting.add(f"AJ{R0}:AJ{RLAST}", FormulaRule(formula=[f'AJ{R0}="Yes"'], fill=FILL_YES))
REG = "'EL scheme register'!"
def rg(col): return f"{REG}${col}${R0}:${col}${RLAST}"

# ---------------- Institution summary
SUM_COLS = [("Institution ID", 11), ("Institution", 24), ("Country", 12), ("Sample group", 18), ("Region group", 18),
    ("QS edition", 8), ("QS rank", 8), ("Review completeness", 12), ("Qualifying current: University-wide", 11),
    ("Qualifying current: Cross-faculty", 11), ("Qualifying current: Faculty/school", 11), ("Qualifying current: Programme/course", 11),
    ("Qualifying current: Unclear reach", 11), ("Qualifying current: total", 11), ("Evidence-incomplete records", 11),
    ("Outside-definition records in register", 11), ("Exclusions / near-misses", 11),
    ("University-wide access evidenced (qualifying current scheme)", 14), ("Compulsory university-wide provision evidenced", 14),
    ("Main qualifying schemes", 46), ("Distribution across faculty areas (qualifying current)", 34),
    ("How embedded the provision appears", 22), ("Concise assessment", 70), ("Material gaps and unresolved questions", 60)]
ws = wb.create_sheet("EL institution summary", 1)
header(ws, "Experiential learning: institution summary",
       "One row per institution in the 124-institution sample. Counts are live formulas over 'EL scheme register' (column AJ = counted, meets all three, current).",
       "Findings within the researched sample, not national prevalence estimates. 'Not started' / 'Partial' institutions must not be read as negative findings. No composite score is used.",
       [c for c, _ in SUM_COLS], [w for _, w in SUM_COLS])
S0 = 7
for n, x in enumerate(summary):
    inst = x["inst"]; rr = S0 + n; r = x["r"] or {}
    mine = [z for z in schemes if z["iid"] == inst["inst_id"]]
    q = [z for z in mine if z["qualifying"]]
    areas = collections.Counter(a for z in q for a in (z["sc"].get("faculty_areas") or []))
    nx = sum(1 for e in exclusions if e[1] == inst["inst_id"])
    if x["status"] == "Not started":
        uw = comp = "Not reviewed"
    else:
        uw = "Yes" if any(z["reach"] == "University-wide" for z in q) else "Not evidenced"
        comp = "Yes" if any(z["reach"] == "University-wide" and z["part"] == "Compulsory for all students" for z in q) else "Not evidenced"
    reach_f = [f'=COUNTIFS({rg("A")},$A{rr},{rg("C")},"{rc}",{rg("AJ")},"Yes")' for rc in REACHES]
    vals = [inst["inst_id"], inst["institution"], inst["country"], inst["cohort"], x["region"], inst["qs_edition"], inst["qs_rank"],
            x["status"], *reach_f, f"=SUM(I{rr}:M{rr})",
            f'=COUNTIFS({rg("A")},$A{rr},{rg("AF")},"Potentially relevant; evidence incomplete",{rg("AI")},"Yes")',
            f'=COUNTIFS({rg("A")},$A{rr},{rg("AF")},"Outside strict definition")',
            f"=COUNTIF('EL exclusions & gaps'!$B$7:$B$5000,$A{rr})", uw, comp,
            s([s(z["sc"].get("official_name")) + (f" ({s(z['sc'].get('module_code'))})" if s(z["sc"].get("module_code")) else "") + f" [{z['reach']}]" for z in q]),
            s([f"{AREA_LABEL.get(a, a)}: {k}" for a, k in areas.most_common()]),
            s(r.get("embeddedness")) or ("Not reviewed" if x["status"] == "Not started" else "Unclear"),
            re.sub(r"\s+", " ", re.sub(r"[#*_`>]", "", s(r.get("profile_md"))))[:1200],
            s(r.get("gaps"), 1500) + (" | PARTIAL REVIEW — remaining: " + s(r.get("partial_remaining")) if x["status"] == "Partial" else "")]
    for j, v in enumerate(vals, 1):
        c = ws.cell(rr, j, v); c.font = F_BODY; c.alignment = WRAP
    ws.cell(rr, 2).font = F_BOLD
NSUM = len(summary); SLAST = S0 + NSUM - 1
add_table(ws, "InstitutionSummary", len(SUM_COLS), NSUM)
dv(ws, f"H{S0}:H{SLAST}", ["Complete", "Partial", "Not started"])
ws.conditional_formatting.add(f"N{S0}:N{SLAST}", FormulaRule(formula=[f"N{S0}>0"], fill=FILL_YES))
ws.conditional_formatting.add(f"H{S0}:H{SLAST}", FormulaRule(formula=[f'H{S0}="Partial"'], fill=FILL_UNC))
ws.conditional_formatting.add(f"H{S0}:H{SLAST}", FormulaRule(formula=[f'H{S0}="Not started"'], fill=FILL_NO))
SUMR = "'EL institution summary'!"
def sr(col): return f"{SUMR}${col}${S0}:${col}${SLAST}"

# ---------------- Exclusions
EX_COLS = [("Exclusion ID", 12), ("Institution ID", 11), ("Institution", 22), ("Region group", 18), ("Scheme / activity", 36),
    ("From legacy research?", 10), ("Legacy row no. (Legacy schemes tab)", 10), ("Category", 26),
    ("Reason it does not qualify / evidence missing", 70), ("Source URL", 40), ("Status note", 30)]
ws = wb.create_sheet("EL exclusions & gaps", 3)
header(ws, "Experiential learning: exclusions, near-misses and evidence gaps",
       "Previously included frameworks and near-misses, with the reason they do not meet the strict definition or cannot yet be verified.",
       "Legacy rows link back to 'Legacy schemes (v3)' (row number in column G; disposition in that tab's column Q).",
       [c for c, _ in EX_COLS], [w for _, w in EX_COLS])
for n, e in enumerate(exclusions):
    for j, v in enumerate(e, 1):
        c = ws.cell(7 + n, j, v); c.font = F_BODY; c.alignment = WRAP
add_table(ws, "Exclusions", len(EX_COLS), len(exclusions))

# ---------------- Sources
SO_COLS = [("Source ID", 16), ("Institution ID", 11), ("Scheme ID", 12), ("Scheme", 30), ("Supports (claim / criterion)", 22),
    ("URL", 50), ("Title", 34), ("Section / subpage / PDF page", 24), ("Publication / academic year", 14), ("Access date", 11),
    ("Retrieval method", 16), ("Supporting extract or precise paraphrase", 80)]
ws = wb.create_sheet("EL sources", 4)
header(ws, "Experiential learning: claim-level sources",
       "Each row is one source supporting named claims for one scheme. 'Supports' lists the criteria/fields it evidences.",
       "Retrieval method 'search-index extract': page text obtained through web-search extracts because direct page fetching was blocked by this environment's network policy. Re-open URLs to confirm before external publication.",
       [c for c, _ in SO_COLS], [w for _, w in SO_COLS])
for n, e in enumerate(sources):
    for j, v in enumerate(e, 1):
        c = ws.cell(7 + n, j, v); c.font = F_BODY; c.alignment = WRAP
add_table(ws, "Sources", len(SO_COLS), len(sources))

# ---------------- Search log
SL_COLS = [("Institution ID", 11), ("Institution", 22), ("Faculty / provision area", 24), ("Search status", 14),
           ("Queries run", 60), ("Key sources examined", 50), ("Finding", 60)]
ws = wb.create_sheet("EL search log", 5)
header(ws, "Experiential learning: institution-level search coverage log",
       "Six areas per institution: central provision, business, engineering/computing, science, arts/humanities/social sciences, other.",
       "Status: Searched / Not applicable (no such faculty) / Not searched (with reason). Institutions not yet researched have no rows.",
       [c for c, _ in SL_COLS], [w for _, w in SL_COLS])
for n, e in enumerate(searchlog):
    for j, v in enumerate(e, 1):
        c = ws.cell(7 + n, j, v); c.font = F_BODY; c.alignment = WRAP
add_table(ws, "SearchLog", len(SL_COLS), len(searchlog))

# ---------------- EL dashboard
ws = wb.create_sheet("EL dashboard", 1)
ws.sheet_view.showGridLines = False
ws["A2"] = "Experiential learning dashboard: collaborative employer-connected projects"; ws["A2"].font = Font(name="Arial", size=18, bold=True, color=NAVY)
ws["A3"] = ("Findings within the researched 124-institution sample (purposive, not random) — not estimates of national prevalence. "
            "A qualifying scheme = project-based AND group working AND connected to a real external organisation (all Yes on the same route) with current delivery confirmed; umbrellas with counted children are not double-counted.")
ws["A3"].font = F_SUB; ws["A3"].alignment = WRAP; ws.merge_cells("A3:L4")
for col, w in zip("ABCDEFGHIJKL", (34, 13, 13, 13, 13, 13, 4, 34, 13, 13, 13, 13)):
    ws.column_dimensions[col].width = w
def kpi(r, c, label, formula, note=""):
    ws.cell(r, c, label).font = F_SUB
    k = ws.cell(r + 1, c, formula); k.font = F_KPI; k.fill = FILL_KPI; k.alignment = Alignment(horizontal="left")
    if note: ws.cell(r + 2, c, note).font = Font(name="Arial", size=8, italic=True, color=SLATE)
kpi(6, 1, "Institutions in scope", f"=COUNTA({sr('A')})")
kpi(6, 3, "Fully reviewed", f'=COUNTIF({sr("H")},"Complete")')
kpi(6, 5, "Partially reviewed", f'=COUNTIF({sr("H")},"Partial")')
kpi(6, 8, "Not yet reviewed", f'=COUNTIF({sr("H")},"Not started")')
kpi(6, 10, "With ≥1 current qualifying scheme", f'=COUNTIF({sr("N")},">0")', "denominator = reviewed (full + partial)")
kpi(10, 1, "Distinct current qualifying schemes", f'=COUNTIF({rg("AJ")},"Yes")')
kpi(10, 3, "Records in scheme register", f"=COUNTA({rg('D')})")
kpi(10, 5, "Exclusion / near-miss records", "=COUNTA('EL exclusions & gaps'!$A$7:$A$5000)")
kpi(10, 8, "Source rows (claim-level)", "=COUNTA('EL sources'!$A$7:$A$20000)")
kpi(10, 10, "Reviewed institutions with none verified", f'=COUNTIFS({sr("H")},"<>Not started",{sr("N")},0)', "\"no qualifying scheme verified through the documented search\"")

def table_block(r0, title, hdrs, rows):
    ws.cell(r0, 1, title).font = F_H2
    for j, h in enumerate(hdrs, 1):
        c = ws.cell(r0 + 1, j, h); c.font = F_HDR; c.fill = FILL_HDR; c.alignment = WRAP
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            c = ws.cell(r0 + 2 + i, j, v); c.font = F_BOLD if j == 1 else F_BODY
    return r0 + 2 + len(rows) + 1

QC_ = rg("AJ")
r0 = 15
r0 = table_block(r0, "Current qualifying schemes by academic reach", ["Academic reach", "Schemes", "Institutions with ≥1", "Share of qualifying schemes"],
    [[rc, f'=COUNTIFS({rg("C")},A{r0 + 2 + i},{QC_},"Yes")', f'=COUNTIF({sr(L(9 + i))},">0")',
      f'=IF(SUM($B${r0 + 2}:$B${r0 + 6})=0,"n/a",B{r0 + 2 + i}/SUM($B${r0 + 2}:$B${r0 + 6}))'] for i, rc in enumerate(REACHES)]
    + [["Total", f"=SUM(B{r0 + 2}:B{r0 + 6})", f'=COUNTIF({sr("N")},">0")', ""]])
for i in range(6): ws.cell(17 + i, 4).number_format = "0%"
r1 = r0
r0 = table_block(r0, "University-wide qualifying schemes: optional access vs compulsory provision",
    ["Participation requirement", "Schemes", "Institutions"],
    [[p, f'=COUNTIFS({rg("C")},"University-wide",{rg("Q")},A{r0 + 2 + i},{QC_},"Yes")',
      f'=SUMPRODUCT(({sr("A")}<>"")/COUNTIF({sr("A")},{sr("A")}&"")*(COUNTIFS({rg("A")},{sr("A")},{rg("C")},"University-wide",{rg("Q")},A{r0 + 2 + i},{QC_},"Yes")>0))']
     for i, p in enumerate(PARTS)])
r0 = table_block(r0, "Current qualifying schemes by external partner type (business-linked vs public/community)",
    ["Partner type", "Schemes", "Institutions"],
    [[p, f'=COUNTIFS({rg("AB")},A{r0 + 2 + i},{QC_},"Yes")',
      f'=SUMPRODUCT(({sr("A")}<>"")/COUNTIF({sr("A")},{sr("A")}&"")*(COUNTIFS({rg("A")},{sr("A")},{rg("AB")},A{r0 + 2 + i},{QC_},"Yes")>0))']
     for i, p in enumerate(PTYPES)])
r0 = table_block(r0, "Evidence status of all register records (counted schemes only; umbrellas excluded)",
    ["Category", "Records"],
    [["Meets all three; current delivery confirmed", f'=COUNTIFS({rg("AF")},"Meets all three",{rg("AG")},"Current delivery confirmed",{rg("AI")},"Yes")'],
     ["Meets all three; current status unclear", f'=COUNTIFS({rg("AF")},"Meets all three",{rg("AG")},"Current status unclear",{rg("AI")},"Yes")'],
     ["Meets all three; historical / paused only", f'=COUNTIFS({rg("AF")},"Meets all three",{rg("AG")},"Historical delivery only",{rg("AI")},"Yes")+COUNTIFS({rg("AF")},"Meets all three",{rg("AG")},"Paused/discontinued",{rg("AI")},"Yes")'],
     ["Meets all three; announced / forthcoming", f'=COUNTIFS({rg("AF")},"Meets all three",{rg("AG")},"Announced/forthcoming",{rg("AI")},"Yes")'],
     ["Potentially relevant; evidence incomplete", f'=COUNTIFS({rg("AF")},"Potentially relevant; evidence incomplete",{rg("AI")},"Yes")'],
     ["Outside strict definition (kept in register as near-miss)", f'=COUNTIFS({rg("AF")},"Outside strict definition",{rg("AI")},"Yes")'],
     ["Excluded / near-miss / evidence-gap records (exclusions tab)", "=COUNTA('EL exclusions & gaps'!$A$7:$A$5000)"],
     ["  of which from legacy research", "=COUNTIF('EL exclusions & gaps'!$F$7:$F$5000,\"Yes\")"]])
# Region breakdown (right-hand block)
ws.cell(15, 8, "Breakdown by country / region group").font = F_H2
rh = ["Region group", "In scope", "Reviewed (full+partial)", "With ≥1 qualifying", "Qualifying schemes", "University-wide qualifying"]
for j, h in enumerate(rh):
    c = ws.cell(16, 8 + j if j == 0 else 8 + j, h); c.font = F_HDR; c.fill = FILL_HDR; c.alignment = WRAP
ws.column_dimensions["M"].width = 13
for i, rgn in enumerate(REGION_ORDER + ["Total"]):
    rr = 17 + i
    ws.cell(rr, 8, rgn).font = F_BOLD
    if rgn == "Total":
        for j in range(5): ws.cell(rr, 9 + j, f"=SUM({L(9 + j)}17:{L(9 + j)}21)").font = F_BOLD
        continue
    ws.cell(rr, 9, f'=COUNTIF({sr("E")},H{rr})')
    ws.cell(rr, 10, f'=COUNTIFS({sr("E")},H{rr},{sr("H")},"<>Not started")')
    ws.cell(rr, 11, f'=COUNTIFS({sr("E")},H{rr},{sr("N")},">0")')
    ws.cell(rr, 12, f'=COUNTIFS({rg("AP")},H{rr},{QC_},"Yes")')
    ws.cell(rr, 13, f'=COUNTIFS({rg("AP")},H{rr},{QC_},"Yes",{rg("C")},"University-wide")')
ws.cell(25, 8, "By country").font = F_H2
for j, h in enumerate(["Country", "In scope", "Reviewed", "With ≥1 qualifying", "Qualifying schemes"]):
    c = ws.cell(26, 8 + j, h); c.font = F_HDR; c.fill = FILL_HDR; c.alignment = WRAP
for i, ctry in enumerate(sorted(set(i_["country"] for i_ in insts), key=lambda z: (z not in ("UK", "Australia", "US"), z))):
    rr = 27 + i
    ws.cell(rr, 8, ctry).font = F_BOLD
    ws.cell(rr, 9, f'=COUNTIF({sr("C")},H{rr})')
    ws.cell(rr, 10, f'=COUNTIFS({sr("C")},H{rr},{sr("H")},"<>Not started")')
    ws.cell(rr, 11, f'=COUNTIFS({sr("C")},H{rr},{sr("N")},">0")')
    ws.cell(rr, 12, f'=COUNTIFS({rg("AO")},H{rr},{QC_},"Yes")')
rlast = max(r0, 27 + 15) + 1
ws.cell(rlast, 1, "Reading the dashboard").font = F_H2
notes = ["Denominators are shown; 'with ≥1 qualifying' should be read against reviewed institutions, not the full 124 if any remain unreviewed.",
         "University-wide = published eligibility open across the institution (subject to stated restrictions). It does not mean every student participates; compulsory university-wide provision is counted separately.",
         "Business-linked schemes are those whose external partners are Business/commercial; Mixed schemes include business and public/charity clients. Public-sector and charity-only schemes are shown separately.",
         "'Evidence incomplete' = no criterion failed, but at least one is Unclear. These are not negative findings.",
         "Evidence was gathered from search-index extracts of official pages (direct page fetching was blocked); see EL method & QC."]
for i, t_ in enumerate(notes):
    c = ws.cell(rlast + 1 + i, 1, "• " + t_); c.font = F_BODY; c.alignment = WRAP; ws.merge_cells(start_row=rlast + 1 + i, start_column=1, end_row=rlast + 1 + i, end_column=12)
    ws.row_dimensions[rlast + 1 + i].height = 28

# Point the original Dashboard EL block at the new review
dash["A56"] = "Experiential-learning research coverage (strict-definition review; detail in 'EL dashboard')"
dash["A57"] = "Institutions in scope / reviewed"; dash["F57"] = f'=COUNTA({sr("A")})'
dash["I57"] = "Current qualifying schemes"; dash["M57"] = f'=COUNTIF({rg("AJ")},"Yes")'
dash["A58"] = "With ≥1 current qualifying scheme"; dash["F58"] = f'=COUNTIF({sr("N")},">0")'
dash["I58"] = "Institutions fully reviewed"; dash["M58"] = f'=COUNTIF({sr("H")},"Complete")'
dash["A60"] = ("Qualifying = project-based, group working and connected to a real external organisation on the same route, with current delivery confirmed. "
               "The legacy scheme sample (152 frameworks/examples) is retained in 'Legacy schemes (v3)'. Lifelong-learning revenue above is NOT revenue from these schemes.")

# ---------------- Method & QC sheet
ws = wb.create_sheet("EL method & QC")
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 120
ws["A2"] = "Experiential learning review: method, definitions, sample reconciliation and quality control"; ws["A2"].font = F_TITLE
rows = [
 ("Research question", "Where do universities provide structured opportunities for students to work collaboratively on substantive projects connected to real employers or businesses, and how widely are these opportunities embedded across each institution?"),
 ("A. Project-based", "Substantive bounded project with a defined problem/brief/research question and identifiable output. Lectures, workshops, talks, visits, networking: insufficient. Placements only with an evidenced project component."),
 ("B. Group working", "Students work together as a team on the project or joint output. Same class, employer mentor, supervisor collaboration or colleagues during an individual placement do not count. Mixed individual/group schemes are split by route or marked Unclear."),
 ("C. Employer / external organisation", "Documented participation of a real external organisation: supplies live brief/problem/data; commissions the work; participates in supervision/mentoring/feedback/assessment; or receives the outputs. Generic industry relevance, hypothetical clients, teaching cases, guest speakers, logos: insufficient. Partner types recorded separately so business-linked provision can be isolated."),
 ("Three-state values", "Yes = explicit supporting evidence; No = evidence establishes the criterion is not met; Unclear = insufficient evidence. Visual fills: green/red/amber; underlying values remain text."),
 ("Classification", "Meets all three = Yes/Yes/Yes. Potentially relevant; evidence incomplete = no No, at least one Unclear. Outside strict definition = at least one No. All three must apply to the same scheme and participation route."),
 ("Academic reach", "University-wide | Cross-faculty | Faculty/school | Programme/course | Unclear — based on published eligibility; administrative host recorded separately."),
 ("Operating status", "Current delivery confirmed | Announced/forthcoming | Historical delivery only | Paused/discontinued | Current status unclear; latest evidenced delivery year recorded separately. A working page is not proof of current delivery."),
 ("Counting", "Distinct scheme or materially different participation route. Umbrella programmes with recorded child routes are flagged and not counted (children are). Annual cohorts, client examples and duplicate webpages are not split."),
 ("Search protocol", "Each institution: central provision + business + engineering/computing + science + arts/humanities/social sciences + other (law/health/design/architecture), using targeted site: searches, module catalogues/handbooks, employer-facing project pages and local-language terms for EU institutions. 'Reviewed' = all applicable areas documented in EL search log; otherwise Partial."),
 ("EVIDENCE LIMITATION", "The research environment's network policy blocked direct page fetching (all hosts denied, Oct 2026). Evidence was therefore obtained through web-search extracts that quote and attribute text to specific official URLs ('search-index extract'). Extracts were not checked against a live copy of each page. Handbook years shown in URLs were recorded; undated pages were not taken as proof of current delivery. Before external publication, re-open source URLs (a spot-check sample at minimum) — enabling network access to university domains in the environment settings would allow this."),
 ("Sample reconciliation", f"Expected 124; found {len(insts)} (Russell Group 24, additional UK 30, Australia 10, US 30, EU 30 across 10 member states). No duplicate institutions; every legacy scheme row maps to an index institution; every institution had ≥1 legacy row. Sample preserved unchanged; no substitutions."),
 ("QS reconciliation", "All rows record QS World University Rankings 2027 edition. 95 ranks cite the ranking table; 29 cite individual QS profile pages (mostly ranks >200) — not re-verified in this review because QS pages could not be opened. Tie flags were internally consistent for ranks shared within the sample. Notes: Adelaide University is the 2026 merger of the University of Adelaide and UniSA; City St George's is the 2024 merger of City and St George's; US/EU rows are not listed in rank order (cosmetic). See Institution index column O."),
 ("Financial data", "Dashboard (lifelong-learning income), UK providers, Revenue evidence tabs and formulas are unchanged; the two Dashboard charts are preserved. Lifelong-learning revenue must not be interpreted as revenue from experiential learning schemes."),
 ("Legacy research", "The 152 legacy rows are retained in 'Legacy schemes (v3)' with a disposition column (Q) pointing to the register or exclusions tab. Institution index columns I–M still count legacy rows."),
]
for i, (k, v) in enumerate(rows):
    ws.cell(4 + i, 1, k).font = F_BOLD; c = ws.cell(4 + i, 2, v); c.font = F_BODY; c.alignment = WRAP
    ws.cell(4 + i, 1).alignment = WRAP
    ws.row_dimensions[4 + i].height = max(15, 15 * (len(v) // 150 + 1))
q0 = 6 + len(rows)
ws.cell(q0, 1, "Automated QC findings (build time)").font = F_H2
errs = [x for x in qc if x[0] == "ERROR"]; warns = [x for x in qc if x[0] == "WARN"]
ws.cell(q0 + 1, 1, "Summary").font = F_BOLD
ws.cell(q0 + 1, 2, f"{len(errs)} errors, {len(warns)} warnings across {len(results)} researched institutions (listed below; resolved items are removed on rebuild).").font = F_BODY
for i, (sev, iid, item, msg) in enumerate(qc[:400]):
    ws.cell(q0 + 2 + i, 1, f"{sev} {iid}").font = F_BODY
    ws.cell(q0 + 2 + i, 2, f"{item}: {msg}").font = F_BODY

order = ["Dashboard", "EL dashboard", "EL institution summary", "EL scheme register", "EL exclusions & gaps", "EL sources",
         "EL search log", "EL method & QC", "Institution index", "Legacy schemes (v3)", "UK providers", "Revenue evidence"]
wb._sheets = [wb[n] for n in order]
wb.active = 1
wb.save(OUT)

# ---------------------------------------------------------------- profiles markdown
md = ["# University experiential learning: institution profiles", "",
      f"Strict-definition review of collaborative, project-based, employer/external-organisation-connected provision. Research date {TODAY}. ",
      "Companion to `University_experiential_learning_detailed_review.xlsx` (scheme IDs refer to the 'EL scheme register' tab). "
      "Findings describe the researched sample only. Evidence was gathered via search-index extracts of official pages because direct page access was blocked; see the workbook's 'EL method & QC' tab.", ""]
for rgn in REGION_ORDER:
    md += [f"## {rgn}", ""]
    for x in summary:
        if x["region"] != rgn: continue
        inst = x["inst"]; r = x["r"]
        md.append(f"### {inst['institution']} ({inst['inst_id']}, {inst['country']})")
        if not r:
            md += ["", "_Not yet reviewed in this pass._", ""]; continue
        mine = [z for z in schemes if z["iid"] == inst["inst_id"]]
        md += ["", f"**Review:** {x['status']}" + (f" — remaining: {s(r.get('partial_remaining'))}" if x["status"] == "Partial" else "") +
               f" · **Embeddedness:** {s(r.get('embeddedness')) or 'Unclear'} · **Current qualifying schemes:** {sum(1 for z in mine if z['qualifying'])}", ""]
        md += [s(r.get("profile_md")), ""]
        if mine:
            md += ["| ID | Scheme | Reach | Participation | A/B/C | Partner | Classification | Status |", "|---|---|---|---|---|---|---|---|"]
            for z in mine:
                nm = s(z["sc"].get("official_name")).replace("|", "/") + (f" ({s(z['sc'].get('module_code'))})" if s(z["sc"].get("module_code")) else "")
                md.append(f"| {z['sid']}{' (umbrella)' if z['umb'] == 'Yes' else ''} | {nm} | {z['reach']} | {z['part']} | {z['a']}/{z['b']}/{z['c']} | {z['ptype']} | {z['cls']} | {z['ost']} {s(z['sc'].get('latest_evidenced_year'))} |")
            md.append("")
        if s(r.get("gaps")): md += [f"**Uncertain / unresolved:** {s(r.get('gaps'))}", ""]
open(PROFILES, "w").write("\n".join(md))

# ---------------------------------------------------------------- checkpoint + QC report
done = [x["inst"]["inst_id"] for x in summary if x["status"] == "Complete"]
part = [x["inst"]["inst_id"] for x in summary if x["status"] == "Partial"]
todo = [x["inst"]["inst_id"] for x in summary if x["status"] == "Not started"]
cp = [f"# Checkpoint — updated {datetime.datetime.now(datetime.timezone.utc):%Y-%m-%d %H:%M} UTC", "",
      f"- Complete: {len(done)} — {', '.join(done)}", f"- Partial: {len(part)} — {', '.join(part)}",
      f"- Not started: {len(todo)} — {', '.join(todo)}", f"- Scheme records: {len(schemes)}; currently qualifying: {sum(z['qualifying'] for z in schemes)}; exclusions: {len(exclusions)}; sources: {len(sources)}",
      "", "## Partial reviews: what remains"]
cp += [f"- {x['inst']['inst_id']} {x['inst']['institution']}: {s((x['r'] or {}).get('partial_remaining'))}" for x in summary if x["status"] == "Partial"]
cp += ["", "## Unresolved points on schemes (evidence incomplete)"]
cp += [f"- {z['sid']} {s(z['sc'].get('official_name'))}: {s(z['sc'].get('unresolved'), 300)}" for z in schemes if z["cls"].startswith("Potentially")]
cp += ["", "Resume: research IDs under 'Not started'/'Partial' following RESEARCH_BRIEF.md, write results/<ID>.json, then run `python3 -I el_review/build.py`."]
open(os.path.join(ELR, "CHECKPOINT.md"), "w").write("\n".join(cp))
open(os.path.join(ELR, "qc_report.txt"), "w").write("\n".join(" | ".join(map(str, q)) for q in qc))
print(f"built: {len(results)} researched, {len(schemes)} schemes, {sum(z['qualifying'] for z in schemes)} qualifying, {len(exclusions)} exclusions, {len(sources)} sources, QC {len(errs)} errors / {len(warns)} warnings")
